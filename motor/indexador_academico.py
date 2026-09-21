#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
Motor de Indexación Modular Académico Universal.
Extrae PDFs hoja a hoja y DOCXs, particionando los índices por materia en
storage_index/{CODIGO_MATERIA}/, compatible con cualquier disciplina
(Informática, Biología, Nutrición Animal, etc.).
"""

import os
import sys
import re
import json
import math
from pathlib import Path
from collections import Counter
import pypdf
import docx

# Importar el auto_gestor local
sys.path.insert(0, str(Path(__file__).resolve().parent))
from auto_gestor import AutoGestor

# Asegurar codificación UTF-8
if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding='utf-8', errors='replace')
        sys.stderr.reconfigure(encoding='utf-8', errors='replace')
    except Exception:
        pass

def limpiar_texto(texto):
    if not texto:
        return ""
    texto = re.sub(r'[ \t]+', ' ', texto)
    texto = re.sub(r'\n{3,}', '\n\n', texto)
    return texto.strip()

def clasificar_tipo_documento(nombre_archivo):
    nombre = nombre_archivo.lower()
    if any(nombre.endswith(ext) for ext in [".png", ".jpg", ".jpeg", ".webp", ".svg", ".gif", ".bmp"]):
        return "recurso_visual"
    if any(k in nombre for k in ["ejercicio", "actividad", "taller", "tarea", "problema"]):
        return "ejercicios"
    if any(k in nombre for k in ["practica", "práctica", "laboratorio", "caso"]):
        return "practica"
    if any(k in nombre for k in ["solucion", "solución", "respuestas"]):
        return "soluciones"
    return "teoria"

class VectorizadorLocal:
    def __init__(self, dim=768, url_lmstudio="http://127.0.0.1:1234/v1/embeddings", modelo_lmstudio="nomic-ai/text-embedding-nomic-embed-text-v1.5"):
        self.dim = dim
        self.url_lmstudio = url_lmstudio
        self.modelo_lmstudio = modelo_lmstudio

    def vectorizar_lmstudio(self, texto):
        """Genera embeddings semánticos profundos mediante la API de LM Studio."""
        import urllib.request
        try:
            payload = json.dumps({
                "model": self.modelo_lmstudio,
                "input": texto
            }).encode('utf-8')
            req = urllib.request.Request(
                self.url_lmstudio,
                data=payload,
                headers={"Content-Type": "application/json"}
            )
            with urllib.request.urlopen(req, timeout=5) as res:
                if res.status == 200:
                    data = json.loads(res.read().decode('utf-8'))
                    return data["data"][0]["embedding"]
        except Exception:
            return None

    def vectorizar(self, texto):
        # 1. Intentar con LM Studio (calidad máxima)
        vec_neuronal = self.vectorizar_lmstudio(texto)
        if vec_neuronal and len(vec_neuronal) > 0:
            return vec_neuronal

        # 2. Fallback local matemático si LM Studio no está disponible
        tokens = re.findall(r'\b[a-zA-ZáéíóúÁÉÍÓÚñÑ0-9_]{2,}\b', texto.lower())
        if not tokens:
            return [0.0] * self.dim
            
        vec = [0.0] * self.dim
        counts = Counter(tokens)
        total = len(tokens)
        
        for token, count in counts.items():
            tf = (count / total) * (1.0 + math.log(1.0 + len(token)))
            h = hash(token) % self.dim
            vec[h] += tf
            if len(token) > 4:
                h_sub = hash(token[:4]) % self.dim
                vec[h_sub] += tf * 0.5
                
        norm = math.sqrt(sum(x * x for x in vec))
        if norm > 0:
            vec = [round(x / norm, 5) for x in vec]
        return vec

class IndexadorAcademico:
    def __init__(self, raiz_proyecto=None):
        self.gestor = AutoGestor(raiz_proyecto)
        self.raiz = self.gestor.raiz
        self.carpeta_storage = self.gestor.carpeta_storage
        self.vectorizador = VectorizadorLocal(dim=768)

    def extraer_pdf(self, ruta_pdf):
        paginas_extraidas = []
        try:
            reader = pypdf.PdfReader(str(ruta_pdf))
            total_paginas = len(reader.pages)
            for num_pag, page in enumerate(reader.pages, start=1):
                texto = page.extract_text() or ""
                texto = limpiar_texto(texto)
                if texto and len(texto) > 30:
                    paginas_extraidas.append({
                        "pagina": num_pag,
                        "total_paginas": total_paginas,
                        "texto": texto
                    })
        except Exception as e:
            print(f"⚠️ Error al leer PDF {ruta_pdf.name}: {e}", file=sys.stderr)
        return paginas_extraidas

    def extraer_docx(self, ruta_docx):
        fragmentos = []
        try:
            doc = docx.Document(str(ruta_docx))
            buffer_texto = []
            num_seccion = 1

            for p in doc.paragraphs:
                txt = limpiar_texto(p.text)
                if txt:
                    buffer_texto.append(txt)
                    if len(" ".join(buffer_texto)) > 600:
                        fragmentos.append({
                            "pagina": num_seccion,
                            "total_paginas": 1,
                            "texto": "\n".join(buffer_texto)
                        })
                        buffer_texto = []
                        num_seccion += 1

            for tabla in doc.tables:
                filas_tabla = []
                for row in tabla.rows:
                    celdas = [limpiar_texto(c.text) for c in row.cells]
                    filas_tabla.append(" | ".join(celdas))
                if filas_tabla:
                    fragmentos.append({
                        "pagina": num_seccion,
                        "total_paginas": 1,
                        "texto": "TABLA:\n" + "\n".join(filas_tabla)
                    })
                    num_seccion += 1

            if buffer_texto:
                fragmentos.append({
                    "pagina": num_seccion,
                    "total_paginas": max(1, num_seccion),
                    "texto": "\n".join(buffer_texto)
                })
        except Exception as e:
            print(f"⚠️ Error al leer DOCX {ruta_docx.name}: {e}", file=sys.stderr)
        return fragmentos

    def extraer_html(self, ruta_html):
        fragmentos = []
        try:
            with open(ruta_html, "r", encoding="utf-8", errors="ignore") as f:
                html_content = f.read()
            # Eliminar scripts y estilos
            limpio = re.sub(r'<script[^>]*>.*?</script>', '', html_content, flags=re.DOTALL | re.IGNORECASE)
            limpio = re.sub(r'<style[^>]*>.*?</style>', '', limpio, flags=re.DOTALL | re.IGNORECASE)
            # Reemplazar saltos de bloque
            limpio = re.sub(r'<(?:h[1-6]|p|div|tr|li)[^>]*>', '\n', limpio, flags=re.IGNORECASE)
            # Eliminar etiquetas restantes
            limpio = re.sub(r'<[^>]+>', ' ', limpio)
            # Normalizar entidades
            limpio = limpio.replace("&nbsp;", " ").replace("&amp;", "&").replace("&lt;", "<").replace("&gt;", ">")
            texto_final = limpiar_texto(limpio)
            if texto_final and len(texto_final) > 30:
                fragmentos.append({
                    "pagina": 1,
                    "total_paginas": 1,
                    "texto": texto_final
                })
        except Exception as e:
            print(f"⚠️ Error al leer HTML {ruta_html.name}: {e}", file=sys.stderr)
        return fragmentos

    def extraer_tex(self, ruta_tex):
        fragmentos = []
        try:
            with open(ruta_tex, "r", encoding="utf-8", errors="ignore") as f:
                raw_tex = f.read()
            sin_comentarios = re.sub(r'(?<!\\)%.*$', '', raw_tex, flags=re.MULTILINE)
            texto_final = limpiar_texto(sin_comentarios)
            if texto_final and len(texto_final) > 20:
                fragmentos.append({
                    "pagina": 1,
                    "total_paginas": 1,
                    "texto": texto_final
                })
        except Exception as e:
            print(f"⚠️ Error al leer TeX {ruta_tex.name}: {e}", file=sys.stderr)
        return fragmentos

    def extraer_imagen(self, ruta_img):
        fragmentos = []
        try:
            from clasificador_pedagogico import extraer_contexto_entorno_imagen
            ctx = extraer_contexto_entorno_imagen(ruta_img)
            if ctx and len(ctx) > 20:
                fragmentos.append({
                    "pagina": 1,
                    "total_paginas": 1,
                    "texto": ctx
                })
        except Exception as e:
            print(f"⚠️ Error al extraer contexto de imagen {ruta_img.name}: {e}", file=sys.stderr)
        return fragmentos

    def fragmentar_texto(self, texto, max_chars=800):
        oraciones = re.split(r'(?<=[.?!])\s+', texto)
        chunks = []
        actual = []
        longitud_actual = 0

        for oracion in oraciones:
            oracion = oracion.strip()
            if not oracion:
                continue
            if longitud_actual + len(oracion) > max_chars and actual:
                chunks.append(" ".join(actual))
                actual = [oracion]
                longitud_actual = len(oracion)
            else:
                actual.append(oracion)
                longitud_actual += len(oracion)

        if actual:
            chunks.append(" ".join(actual))
        return chunks or [texto]

    def indexar_todo(self):
        revision = self.gestor.verificar_proyecto()
        if revision["estado"] == "PROYECTO_VACIO":
            return {
                "exito": False,
                "mensaje": revision["mensaje_alumno"]
            }

        materias = revision.get("materias", {})
        resumen = {
            "materias": {},
            "total_fragmentos": 0,
            "total_documentos": revision.get("total_archivos", 0)
        }

        self.carpeta_storage.mkdir(parents=True, exist_ok=True)

        EXTS_INDEXABLES = (
            "*.pdf", "*.docx", "*.html", "*.txt", "*.md", "*.tex",
            "*.png", "*.jpg", "*.jpeg", "*.webp", "*.svg", "*.gif", "*.bmp"
        )

        for codigo, info in materias.items():
            print(f"📚 Indexando Materia: {codigo} - {info['nombre']}...")
            ruta_materia = Path(info["ruta"])
            carpeta_indice = self.carpeta_storage / codigo
            carpeta_indice.mkdir(parents=True, exist_ok=True)

            if ruta_materia.is_file():
                documentos = [ruta_materia]
            else:
                documentos = []
                for ext in EXTS_INDEXABLES:
                    documentos.extend(ruta_materia.rglob(ext))
                documentos = [
                    d for d in documentos 
                    if not any(part.startswith('.') for part in d.parts)
                    and d.name.lower() not in {"requirements.txt", "system_prompt_tutor.txt", "temario_activo.md"}
                ]
            nodos_materia = []
            temas_detectados = set()

            for doc in documentos:
                nombre_archivo = doc.name
                tipo_doc = clasificar_tipo_documento(nombre_archivo)
                try:
                    rel_padre = doc.parent.relative_to(ruta_materia)
                    tema_nombre = str(rel_padre).split(os.sep)[0] if str(rel_padre) != "." else "General"
                except Exception:
                    tema_nombre = "General"
                temas_detectados.add(tema_nombre)

                suf = doc.suffix.lower()
                if suf == ".pdf":
                    paginas = self.extraer_pdf(doc)
                elif suf == ".docx":
                    paginas = self.extraer_docx(doc)
                elif suf in (".html", ".htm"):
                    paginas = self.extraer_html(doc)
                elif suf == ".tex":
                    paginas = self.extraer_tex(doc)
                elif suf in (".png", ".jpg", ".jpeg", ".webp", ".svg", ".gif", ".bmp"):
                    paginas = self.extraer_imagen(doc)
                else:
                    try:
                        with open(doc, "r", encoding="utf-8", errors="ignore") as f:
                            t = limpiar_texto(f.read())
                        paginas = [{"pagina": 1, "total_paginas": 1, "texto": t}] if t else []
                    except Exception:
                        paginas = []

                for p in paginas:
                    num_pag = p["pagina"]
                    chunks = self.fragmentar_texto(p["texto"])
                    for idx_chunk, chunk in enumerate(chunks):
                        nodo_id = f"{codigo}_{re.sub(r'[^a-zA-Z0-9]', '_', tema_nombre)}_p{num_pag}_c{idx_chunk}"
                        vector = self.vectorizador.vectorizar(chunk)
                        
                        nodo = {
                            "id": nodo_id,
                            "materia_codigo": codigo,
                            "materia_nombre": info["nombre"],
                            "tema": tema_nombre,
                            "archivo": nombre_archivo,
                            "pagina": num_pag,
                            "tipo": tipo_doc,
                            "texto": chunk,
                            "vector": vector
                        }
                        nodos_materia.append(nodo)

            ruta_indice_json = carpeta_indice / "index.json"
            with open(ruta_indice_json, "w", encoding="utf-8") as f:
                json.dump({
                    "materia_codigo": codigo,
                    "nombre": info["nombre"],
                    "temas": sorted(list(temas_detectados)),
                    "total_nodos": len(nodos_materia),
                    "nodos": nodos_materia
                }, f, indent=2, ensure_ascii=False)

            resumen["materias"][codigo] = {
                "nombre": info["nombre"],
                "temas": sorted(list(temas_detectados)),
                "documentos": len(documentos),
                "fragmentos_indexados": len(nodos_materia)
            }
            resumen["total_fragmentos"] += len(nodos_materia)

        self.gestor.actualizar_manifest()
        return resumen

if __name__ == "__main__":
    indexador = IndexadorAcademico()
    resultado = indexador.indexar_todo()
    print(json.dumps(resultado, indent=2, ensure_ascii=False))
